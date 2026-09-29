#!/usr/bin/env python3
"""Build editable defense PPTX and vector PDF from one declarative layout.

Reading saved measurements and existing scientific figures only. No simulation,
surface evaluation, Office rendering, or modification of original photographs.
Run with PYTHONPATH=tmp/final-writing-python:tmp/thesis-python_20260928/site-packages
  python3 scripts/build_final_defense_20260929.py --output docs/thesis/defense_20260929
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import math
import statistics
import tempfile
import textwrap
import zipfile
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.font_manager import FontProperties
from matplotlib.patches import Rectangle
from PIL import Image, ImageDraw, ImageFont
import pypdfium2 as pdfium
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
W, H = 13.333333, 7.5
NAVY, TEAL, GRAY, PALE, ORANGE = '#17324D', '#087F8C', '#536373', '#EDF4F7', '#C97939'
BLUE, RED = '#4878A6', '#B35353'
FONT = Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
BOLD = Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc')
FONTNAME = 'Noto Sans CJK SC'
FIG = Path('docs/thesis/figures')
INPUTS: set[Path] = set()


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def source(path):
    p = Path(path)
    p = p if p.is_absolute() else ROOT / p
    if not p.is_file():
        raise FileNotFoundError(p)
    INPUTS.add(p)
    return p


def read_json(path):
    return json.loads(source(path).read_text())


def wrapped(text, width, size, bold=False):
    font = ImageFont.truetype(str(BOLD if bold else FONT), round(size * 4))
    limit = (width * 72 - 2) * 4
    lines = []
    for paragraph in text.split('\n'):
        current = ''
        for char in paragraph:
            if current and font.getlength(current + char) > limit:
                lines.append(current.rstrip())
                current = char.lstrip()
            else:
                current += char
        lines.append(current)
    return '\n'.join(lines)


def add_text(s, text, x, y, w, size=20, color=NAVY, bold=False, align='left', line=1.30):
    text = wrapped(text.replace('₀', '0').replace('₅', '5'), w, size, bold)
    h = (text.count('\n') + 1) * size * line / 72 + .05
    if y + h > 7.05:
        raise ValueError(f'text below content boundary: {s["title"]}: {text}')
    s['elements'].append(dict(kind='text', text=text, x=x, y=y, w=w, h=h,
                              size=size, color=color, bold=bold, align=align, line=line))
    return y + h


def rect(s, x, y, w, h, color=PALE, edge=None):
    s['elements'].append(dict(kind='rect', x=x, y=y, w=w, h=h, color=color, edge=edge))


def picture(s, stem, x, y, w, h):
    p = source(str(stem) + '.png')
    im = Image.open(p)
    scale = min(w / im.width, h / im.height)
    iw, ih = im.width * scale, im.height * scale
    paths = {'png': str(p.relative_to(ROOT))}
    for ext in ('pdf', 'svg'):
        q = ROOT / (str(stem) + '.' + ext)
        if q.exists():
            source(q)
            paths[ext] = str(q.relative_to(ROOT))
    s['elements'].append(dict(kind='picture', x=x + (w-iw)/2, y=y+(h-ih)/2,
                              w=iw, h=ih, paths=paths))


def photo(s, path, x, y, w, h):
    p = source(path)
    im = Image.open(p)
    scale = min(w/im.width, h/im.height)
    iw, ih = im.width*scale, im.height*scale
    s['elements'].append(dict(kind='picture', x=x+(w-iw)/2, y=y+(h-ih)/2,
                              w=iw, h=ih, paths={'png': str(p.relative_to(ROOT))}))


def slide(title, kicker, notes):
    s = dict(title=title, kicker=kicker, notes=notes, elements=[])
    rect(s, 0, 0, W, .13, TEAL)
    add_text(s, kicker, .55, .32, 12.2, 12, TEAL, True)
    add_text(s, title, .55, .69, 12.2, 27, NAVY, True)
    return s


def bullets(s, items, x=.65, y=1.7, w=11.95, size=22, gap=.25):
    for item in items:
        rect(s, x, y+.12, .07, .20, TEAL)
        y = add_text(s, item, x+.22, y, w-.22, size) + gap
    return y


def note(s, text):
    add_text(s, text, .62, 6.49, 12.08, 11.8, GRAY)


def table(s, headers, rows, widths, x=.65, y=1.68, rowh=.61, size=17):
    total = sum(widths)
    for r, values in enumerate([headers] + rows):
        yy = y + r*rowh
        rect(s, x, yy, total, rowh, NAVY if r == 0 else ('#F0F5F8' if r%2 else '#FFFFFF'))
        xx = x
        for val, width in zip(values, widths):
            add_text(s, str(val), xx+.13, yy+.13, width-.24, size,
                     '#FFFFFF' if r == 0 else NAVY, r == 0)
            xx += width


def metric_card(s, value, label, x, y=1.70, w=3.80, h=1.25, color=TEAL):
    rect(s, x, y, w, h)
    add_text(s, value, x+.18, y+.15, w-.36, 30, color, True)
    add_text(s, label, x+.18, y+.77, w-.36, 14, GRAY)


def collect_data():
    v36 = read_json('audit_results/v36_online_confirmation_20260918/result.json')
    ext = read_json('audit_results/thesis_expansion_20260928/analysis_review/summary.json')
    final = read_json('audit_results/final_local_layout_validation_20260929/analysis/summary.json')
    root = 'audit_results/final_local_layout_validation_20260929/analysis/'
    rows = list(csv.DictReader(source(root+'episode_metrics.csv').open()))
    read_json(root+'manifest.json')
    read_json(root+'paired_metrics.json')
    if final['declared_tasks'] != 8 or final['completed'] + final['failed'] != 8:
        raise ValueError('Final 8-task inventory is not terminal; final defense cannot be built.')
    if len(rows) != 8 or len({r['id'] for r in rows}) != 8:
        raise ValueError('Final results must contain all eight unique tasks.')
    means = {m: statistics.mean(float(r['05cm_joint']) for r in rows if r['mode']==m)
             for m in ['G', 'S']}
    pairs = []
    for scene in sorted({r['scene'] for r in rows}):
        for h in [0, 1]:
            p = {r['mode']:r for r in rows if r['scene']==scene and int(r['hypothesis'])==h}
            if set(p) != {'G','S'}:
                raise ValueError('Missing method arm')
            pairs.append(dict(scene=scene, h=h, G=float(p['G']['05cm_joint']),
                              S=float(p['S']['05cm_joint'])))
    rows36 = read_json('audit_results/article_stage_20260928/analysis_v1/article36_final_summary_v1/rows36.json')
    if len(rows36) != 36:
        raise ValueError('Expected complete 36-task table')
    return dict(v36=v36, ext=ext, final=final, rows=rows, means=means, pairs=pairs,
                new_relative_percent=100*(means['S']/means['G']-1), rows36=rows36)


def make_slides(d):
    slides = []
    def new(*args):
        s = slide(*args); slides.append(s); return s

    s = new('面向工业设施的语义—几何主动三维建图', '毕业论文答辩 · 2026.09.29',
        '本研究面向设备布置调整后的环境模型获取，讨论有限预算下如何选择观察方向。研究主体是在CPU虚拟环境中完成的闭环实验。这里的小车照片只说明后续接口条件，不作为实车效果证据。学校、学位、专业及姓名等信息按要求留空。')
    add_text(s, '预算受限的主动观测\n类别先验 · 几何反馈 · 实测重建', .68, 1.85, 7.2, 29, TEAL, True)
    add_text(s, '学校：____________    学位：____________\n专业：____________\n答辩人：__________    指导教师：__________', .70, 4.28, 7.5, 18, GRAY)
    photo(s, 'docs/thesis/assets/robot_20260928/robot_cutout_user.png', 8.6, 1.72, 4.02, 4.65)
    note(s, '图示为用户提供的原始机器人照片；本报告中的性能结果均来自虚拟环境。')

    s = new('工业需求：走遍通道，还不等于建好设备表面', '01  问题与研究目标',
        '工业生产中设备和工位的布置可能在任务之间发生变化，机器人因此需要重新获取地图。本研究只处理一次任务内静态环境的采集，不研究动态避障或跨时变化检测。有限预算下，从哪一侧看设备会影响遮挡面恢复；因此需要把二维覆盖和三维表面质量同时纳入分析。')
    for x, n, title, desc in [(0.65,'01','布局调整','设备与工位布置改变\n需要获取当前环境模型'),(4.78,'02','观察有代价','行驶、转向均消耗预算\n并须保留完整返航余量'),(8.91,'03','表面有遮挡','相似初始轮廓可能对应\n不同的有效观察方向')]:
        rect(s,x,1.83,3.77,2.75)
        add_text(s,n,x+.20,2.02,3.36,28,TEAL,True)
        add_text(s,title,x+.20,2.65,3.36,23,NAVY,True)
        add_text(s,desc,x+.20,3.25,3.36,19,GRAY)
    add_text(s,'研究问题：类别线索能否在几何尚有歧义时，\n帮助有限预算内的观测选择，并改善实测表面？',.77,5.08,11.9,25,NAVY,True)
    note(s,'研究范围：单次静态采集；工业布局变化是需求背景，不等同于已实现变化检测。')

    s = new('三条贡献：方法、机制与可复核的条件性证据', '02  贡献定位',
        '本文贡献是面向该任务的方法组合和机制证据，而不是把贝叶斯更新、乘积指标或四个名称分别当作原创算法。第一，将有限类别先验与主动诊断、预算和返航约束连接。第二，解释类别何时改变方向、错误线索如何被实测几何纠正。第三，用实际路径、深度融合网格及完整预算边界验证条件性收益。')
    bullets(s, ['方法：将可纠正的类别—构型先验、主动诊断、\n返航约束与实测建图连接为一个执行闭环。',
                '机制：用简化决策分析和保存的决策记录解释\n“提前选向”与“错误先验纠正”。',
                '证据：联合路径、实际网格、重建检查点与预算边界，\n说明收益成立、持平及消失的条件。'], y=1.82,size=23,gap=.35)
    note(s,'定位为任务化方法与机制贡献；不把标准贝叶斯更新、乘积指标本身或四接口命名视为新算法。')

    s = new('强几何对照：G 同样能够诊断、补看与纠错', '03  局部任务与公平对照',
        '为了隔离类别信息，G不是只向前走的弱基线。G和S共享两个公开模板、安全导航图、主动诊断、重访和同一建图器，主要差别是类别先验是否启用。局部任务预算为42个原子动作，共同执行18个动作前缀，再由策略选向。X与Xnf用于错误先验分析，但Xnf同时关闭实测纠正和未来诊断预测，因此是完整策略对照。')
    table(s,['策略','类别先验','几何诊断 / 实测纠正'],[
        ['G','关闭；初始均匀','保留主动诊断、补看和反馈'],
        ['S','正确关联的有限先验','与 G 相同'],
        ['X','类别关联调换','保留'],
        ['Xnf','类别关联调换','关闭实测纠正及未来诊断预测']], [1.2,3.8,7.0],size=18)
    bullets(s,['同图、同模板、同噪声设定、同 TSDF；质量只在离线评价。',
               '局部 B = 42；共同前缀 18；覆盖门槛 80%；返回起始位置与朝向。'],y=5.06,size=20,gap=.15)
    note(s,'这里的局部 S 表示“启用类别”；后续场景级 S 表示“额外跨设施共享”，两者不能混为一谈。')

    s = new('四接口与两层闭环：预测决定去哪里，实测形成地图', '04  实际执行的系统架构',
        '图示与局部CPU实现对应。OV-SDF登记类别证据，IGCR将当前深度与扫描和模板比较，二者共同形成一次信念更新。STGHP用信念、曝光状态和预算预测下一动作。RPN-UQ检查可执行性及全姿态返航。二维占据与三维TSDF走独立实测分支，预测的表面不会填入实际网格。底部实拍平台只是后续ROS接口示意。')
    picture(s,FIG/'final_architecture_20260929/final_architecture',.55,1.46,5.6,5.40)
    bullets(s,['OV-SDF：类别 → 有限构型先验',
               'IGCR：实测几何 → 纠正证据',
               'STGHP：信念 + 预算 → 下一动作',
               'RPN-UQ：合法性 + 全姿态返航'],x=6.52,y=1.92,w=6.0,size=20,gap=.23)
    rect(s,6.56,5.04,6.06,1.21)
    add_text(s,'实际深度 → TSDF → 封存 → 离线评价\n评价真值不反馈给规划器。',6.75,5.23,5.7,19,TEAL,True)

    s = new('理论解释：类别可信度、诊断代价与方向价值共同作用', '05  信念与决策模型',
        '局部任务用两种公开构型形成有限假设空间。类别与几何以对数证据组合，类别不会成为不可推翻的真值。右侧简化模型假设两种构型、两条可返航路线和对称诊断可靠度。诊断的价值取决于它比当前信心能多提高多少方向正确率，并扣除代价。模型解释一种机制，不保证实际TSDF质量单调上升，也不是新的贝叶斯定理。')
    rect(s,.65,1.78,5.82,4.44); rect(s,6.72,1.78,5.95,4.44)
    add_text(s,'实测信念更新',.91,2.02,5.3,24,TEAL,True)
    add_text(s,'p(h₀) = σ(Lg + Ls)',.91,2.65,5.25,29,NAVY,True)
    add_text(s,'Ls：类别对数证据\nLg：深度 / 扫描的几何证据\n类别初始强度 0.9，可被实测推翻。',.91,3.45,5.2,20)
    add_text(s,'预测中的 0.99 为固定诊断参数，\n不是测得的传感器准确率。',.91,5.06,5.2,16,GRAY)
    add_text(s,'一次诊断的净价值',6.98,2.02,5.4,24,TEAL,True)
    add_text(s,'Δ · max(0, rd − m) − κ',6.98,2.65,5.4,27,NAVY,True)
    add_text(s,'m = max(p, 1 − p)\nΔ：匹配与不匹配方向的价值差\nrd：诊断可靠度；κ：诊断代价',6.98,3.44,5.35,20)
    add_text(s,'几何已明确、方向等价或诊断很便宜时，\n类别的额外价值可能消失。',6.98,5.05,5.38,16,GRAY)
    note(s,'简化模型还假设诊断后两方向均可行、条件独立和对称可靠度；用于机制解释而非全局最优证明。')

    s = new('规划与安全执行：一步落实，保留返航', '06  STGHP / RPN-UQ',
        '每次规划最多向前预测一个未来诊断事件，比较两分支的期望后续价值，但只执行第一个原子动作。取得新观测后重新规划，不会把预测信念当成实测信念。局部原子动作是前进一米或左右转九十度，每次动作后产生付费观测。执行层检查动作合法、当前剩余预算足够以及最终位置和朝向都能回到起点。该保证依赖已给定的安全图和动作模型。')
    for x, lab, desc in [(0.65,'1  预测','至多一次未来诊断\n比较可执行后续价值'),(4.78,'2  执行','只下发第一原子动作\n前进 1 m / 转向 90°'),(8.91,'3  重规划','获取实际 RGB-D / 扫描\n更新证据后再决策')]:
        rect(s,x,1.83,3.77,2.3)
        add_text(s,lab,x+.2,2.03,3.37,23,TEAL,True)
        add_text(s,desc,x+.2,2.73,3.37,18)
    add_text(s,'1 + d返回(x下一步, 起始全姿态) ≤ b剩余',.87,4.70,11.8,30,NAVY,True,align='center')
    add_text(s,'返航检查约束可执行性；并不承诺全局最优或质量每步上升。',.9,5.56,11.5,22,GRAY,align='center')
    note(s,'局部任务没有额外 observe 原子动作；每个执行动作后观测一次，另有初始零动作帧。')

    s = new('评价：既要覆盖，也要真实表面的精度与召回', '07  指标与信息边界',
        '主指标是二维已知覆盖与五厘米阈值表面F1的乘积，同时分列覆盖、精度和召回。精度衡量预测表面与真值的接近程度，召回衡量固定完整外表面被恢复的比例。局部评价使用两模板共同定义的ROI并去除地面，不按方法单独裁剪。场景级扩展使用全网格和不同覆盖定义，不能与局部J直接混合。实验采用实际深度融合，但使用准确位姿，并未验证定位漂移改进。')
    metric_card(s,'Cmap','可达栅格中已知部分的比例',.65)
    metric_card(s,'F1@5 cm','表面精度与召回的调和平均',4.78)
    metric_card(s,'J₅ = Cmap × F1','主指标；P / R / C 同时保留',8.91)
    bullets(s,['相同 ROI：两模板公共范围；固定完整外表面参考，统一去地面。',
               '实际 RGB-D / 扫描与位姿进入建图；模板预测不补进 TSDF。',
               '精确位姿与公开安全图是输入条件；不声称提高 SLAM 定位精度。'],y=3.43,size=21,gap=.29)
    note(s,'局部 Cmap 包含已知占据与空闲格；场景级 Cnav 只计实测可导航空闲覆盖，Jnav 使用全网格评价。')

    s = new('原确认实验：平均 J₅ 提升 6.16%，2 胜 2 平', '08  局部类别作用 · 8 条轨迹 / 4 对条件',
        '原确认在两个已见父布局、各两构型上使用冻结控制器和保留噪声种子，共八条真实闭环轨迹。平均J从0.702576提高到0.745881，即相对6.1638%。四个条件中两胜两平，覆盖完全相同，因此收益主要来自表面质量差异。布局并不是全新类别，这组证据支持受控任务中的类别作用，不代表普遍统计结论。')
    old = d['v36']['overall_comparison']['05cm']
    metric_card(s,f"{old['mean_joint']['G']:.4f} → {old['mean_joint']['S']:.4f}",'平均 J₅：G → S',.65,w=4.5)
    metric_card(s,f"+{old['S_vs_G']['relative_mean_joint_difference']*100:.2f}%",'相对均值增益',5.37,w=3.22)
    metric_card(s,'2 胜 / 2 平','4 个配对条件；全部完成',8.81,w=3.86)
    pairrows=[]
    for parent in ['P00','P01']:
        for h in [0,1]:
            t={r['mode']:r['stages']['final']['measurement']['05cm']['joint'] for r in d['v36']['table'] if r['parent']==parent and r['hypothesis']==h}
            pairrows.append([f'{parent} / h{h}',f"{t['G']:.4f}",f"{t['S']:.4f}",f"{t['S']-t['G']:+.4f}"])
    table(s,['条件','G','S','ΔJ₅'],pairrows,[3,3,3,3],y=3.30,rowh=.55,size=18)
    note(s,'相同平均覆盖 0.952544；均为 42 个付费动作、无碰撞、完整返航。独立布局单位为 2。')

    s = new('路径证据：类别首先改变第 19 步的观察方向', '09  真实路径与动作成本',
        '图中每一对使用相同坐标范围，并明确起点、方向和转向。构型h1在前18个动作及非语义观测相同时，第19步首先分歧，因此能把首次选择差异与类别证据联系起来。P00的h1中G平移28米而S为26米；P01中两者都为28米。这说明质量收益不是简单依靠更多路程，方向分配也很重要。后续所有路径都会重新受几何反馈影响。')
    picture(s,FIG/'article_trajectories_20260928/local_paired_trajectories',.61,1.49,6.15,5.45)
    bullets(s,['同一坐标尺度、相同起点与前缀。',
               'h1：第 19 个执行动作首次分歧。',
               'P00 / h1：G 28 m，S 26 m。',
               'P01 / h1：两者均为 28 m。'],x=7.02,y=1.96,w=5.63,size=21,gap=.26)
    add_text(s,'解释重点：有限预算如何分配给\n不同观察方向，而非“走得越多越好”。',7.10,5.05,5.52,21,TEAL,True)

    s = new('三维证据：展示实际融合网格，而非预测模板', '10  同视角、同尺度的表面恢复',
        '这里是保存深度实际融合得到的网格，使用统一相机和尺度展示，并保留完整四个局部配对条件。用于规划的模板只预测潜在可见表面，不会直接加入TSDF。读图时结合数值精度和召回，不把展示颜色理解成误差，也不通过填孔或美化去制造完整性。变化集中在方向选择影响的侧面和遮挡区域，构型h0则保持一致。')
    picture(s,FIG/'scene_details_20260928/mesh_comparison',.60,1.46,6.2,5.53)
    bullets(s,['所有方法使用相同融合后端。',
               '相同视角、世界尺度与显示规则。',
               '四个配对条件全部展示。',
               '结合 P / R / F1 判断恢复质量。'],x=7.05,y=1.96,w=5.55,size=21,gap=.24)
    rect(s,7.09,5.00,5.51,1.22)
    add_text(s,'网格来自测量；\n预测的可见面积不等于已重建面积。',7.27,5.18,5.1,20,TEAL,True)

    s = new('错误先验能够被纠正：完整策略收益 +2.62%', '11  IGCR 反馈与动作变化',
        '错误类别关联把真实构型的初始概率压到0.1。第28步获取有辨识力的几何后，概率变为0.978178，第29步纠错策略转为前进，而关闭纠错的策略仍转向。完整X与Xnf对照平均提升2.6241%。需要强调Xnf同时去除了未来诊断预测，不能把该收益完全归因于一条后验公式。保存状态上的两次后验干预只支持下一动作变化，不能虚构反事实终点收益。')
    picture(s,FIG/'virtual_method_20260928/publication_timeline',.74,1.65,11.87,3.39)
    add_text(s,'第 28 步：真实构型概率 0.10 → 0.9782\n第 29 步：反馈改变下一动作',.82,5.31,7.12,22,TEAL,True)
    add_text(s,'平均 J₅\n0.655065 → 0.672254',8.58,5.33,4.05,21,NAVY,True)
    note(s,'X/Xnf 是完整纠错策略对照；2 cm 下仍存在一个局部负差，未被删去。')

    s = new('预算边界：语义收益随可用预算改变', '12  固定 128 次扩展 · 保留所有失败与负结果',
        '扩展试验一共128次，96次达到资格条件，32次低预算规划失败。B为30时共同18步前缀已经消耗预算，当前规划不能完成双模板覆盖和返航；这不是所有可能算法都不可行的物理证明。中间预算时原布局和同族变体都有正收益，而B为54时G已经有足够机会获得几何，S的平均差反转。该边界与类别价值是有条件的主张一致。')
    metric_card(s,'128 次','固定清单；不同噪声不算新布局',.65)
    metric_card(s,'96 合格','全部合格任务纳入质量汇总',4.78)
    metric_card(s,'32 失败','B = 30；保留记录与失败分母',8.91,color=ORANGE)
    extrows=[]
    for family,b,lab in [('originals',42,'原布局 / B42'),('same_family_variants',42,'同族变体 / B42'),('originals',54,'原布局 / B54')]:
        a=next(v for v in d['ext']['aggregate'] if v['scene_family']==family and v['budget']==b and v['treatment']=='S' and v['control']=='G')
        extrows.append([lab,f"{a['mean_J5_control']:.4f}",f"{a['mean_J5_treatment']:.4f}",f"{a['relative_percent']:+.2f}%"])
    table(s,['条件','G 均值','S 均值','相对增益'],extrows,[4.5,2.5,2.5,2.5],y=3.41,rowh=.63,size=18)
    add_text(s,'高预算下的负差保留：类别先验不是始终有益的“额外奖励”。',.84,6.01,11.7,22,TEAL,True)

    s = new('40 个实测检查点：质量并不随每一步单调增加', '13  原 8 条轨迹的离线重融合',
        '使用原八条轨迹中保存的传感数据，在18、24、30、36、42步共重融合40次，没有重跑策略。配对方法在五个时刻的覆盖均相同，30步时已经达到终点覆盖，但表面质量仍继续变化。P00的h1在24步出现微小负差，到终点转为正差。检查点是相同轨迹的相关测量，不是40个独立实验，也不把中途状态当成已完成返航的合格终点。')
    picture(s,FIG/'article_reconstruction_submission_20260928_final/checkpoint_quality_submission',.64,1.52,6.12,5.36)
    bullets(s,['8 条保存轨迹 × 5 个固定时刻。',
               '18 / 24 / 30 / 36 / 42 个动作。',
               '30 步覆盖已达终点，质量仍变化。',
               'P00 / h1 第 24 步保留短时负差。'],x=7.02,y=1.94,w=5.61,size=21,gap=.27)
    add_text(s,'这些是离线重融合测量，\n不是 40 次新的在线实验。',7.1,5.28,5.44,21,TEAL,True)

    s = new('独立布局补证：8 条全部完成，1 胜 3 平', '14  新增两种同族背景布局 · 冻结控制器',
        f"最后一轮预先固定两个新背景布局、两个构型、G和S，共八条任务，控制器和评价保持不变。全部任务完成、无碰撞并返航。平均J从{d['means']['G']:.6f}到{d['means']['S']:.6f}，相对提升{d['new_relative_percent']:.4f}%。四个条件中只有后隔断h1为正，其余三组持平；首次动作分歧仍是19。这里新增的是背景和导航布局，两个设备模板已公开，不是新类别泛化。预留额外八次任务未使用。")
    metric_card(s,f"+{d['new_relative_percent']:.2f}%",'新布局平均 J₅ 相对增益',.65,w=3.77)
    metric_card(s,'1 胜 / 3 平','全部四个配对条件',4.78,w=3.77)
    metric_card(s,f"{d['final']['qualified']} / 8 合格",'无新增重试；剩余名额不使用',8.91,w=3.77)
    labels={'L02_side_column':'侧柱布局','L03_rear_partition':'后隔断布局'}
    table(s,['条件','G','S','ΔJ₅'],[[f"{labels[p['scene']]} / h{p['h']}",f"{p['G']:.4f}",f"{p['S']:.4f}",f"{p['S']-p['G']:+.4f}"] for p in d['pairs']],[4.5,2.5,2.5,2.5],y=3.32,rowh=.55,size=18)
    note(s,'新布局单位 = 2；沿用公开双模板与类别映射。与原 +6.16% 分开报告，不宣称新类别泛化。')

    s = new('新增场景展示：同背景对照，真实路径与网格', '15  新布局 / 后隔断 h1 · 与主表一致',
        '左图展示新增布局下的完整配对路径，右图为实际深度融合的表面。所有条件都保留在源图与数据中。后隔断h1是唯一动作和结果产生差异的配对：G与S保持相同非语义前缀，S在19步改向，终点联合指标提高0.084109。侧柱两个构型及后隔断h0持平，因此我们把结论限定为方向选择仍有余地的条件，不把新背景个数包装成大量新类别。')
    picture(s,FIG/'final_validation_20260929/transfer_paths',.58,1.48,6.12,4.85)
    picture(s,FIG/'final_validation_20260929/transfer_meshes',6.9,1.48,5.86,4.85)
    note(s,'全部配对条件展示；几何与坐标尺度由原绘图协议固定。改进并未在每个新布局出现。')

    s = new('36 次场景级扩展：共享持平，不扩大语义主张', '16  多设施系统验证与适用边界',
        '场景级扩展包括原版、地面关联修正和曝光去重修正，共三版、三个布局、四种策略，合计36次。局部S和这里的S不是同一干预：这里B已具有类别先验，S只增加跨实例共享。全部九组S/B的实际动作和终点质量相同。原版一个CELL/G任务运动完成但原评价失败，统一离线派生质量不会改变其原失败身份。最新版本还显示类别有负例、诊断策略有正例，均不能包装成共享创新被验证。')
    rows=[]
    for scene,lab in [('AISLE','巷道'),('CELL','工位'),('LOOP','环路')]:
        rr={r['method']:r for r in d['rows36'] if r['arm']=='ExposureV3' and scene in r['scene_id']}
        rows.append([lab,f"{rr['G']['J_nav']:.4f}",f"{rr['B']['J_nav']:.4f}",f"{rr['S']['J_nav']:.4f}",f"{rr['NBV']['J_nav']:.4f}"])
    table(s,['最终版 / Jnav','G','B','S','NBV'],rows,[4,2,2,2,2],y=1.80,rowh=.67,size=19)
    bullets(s,['9 组 S/B：实际动作与质量全相同，未验证额外共享收益。',
               '局部 S/G 检验类别可用性；这里 S/B 只检验额外共享。',
               '原流程 35 / 36 合格；1 条原评价失败保留，另有派生质量。'],y=4.80,size=19,gap=.10)
    note(s,'场景级 Jnav 为全网格、多设施宏平均质量 × 实测空闲覆盖；不能与局部 J₅ 数值直接合并。')

    s = new('结论边界清楚，才能说明方案在哪些任务中有价值', '17  讨论与后续应用',
        '结果支持的是：在几何尚有歧义、方向有价值差、预算有限时，类别先验可以改变早期观察并改善实际表面。几何反馈为错误先验提供纠正机会。当前局限包括有限构型模板、合成类别线索、准确位姿和公开安全图；高预算与共享实验也给出了明确边界。ROS1和ZED接口可用于未来迁移，但本论文不依赖额外实车性能才能完成仿真结论。')
    bullets(s,['已支持：受控条件下的类别提前选向；完整策略中的错误先验纠正。',
               '已保留：高预算负差、低预算失败、新布局持平与共享无增益。',
               '范围：有限模板、合成类别线索、精确位姿、公开安全图。',
               '后续：自然语义、定位误差与 ROS 1 平台验证，作为迁移补充。'],y=1.83,size=23,gap=.29)
    note(s,'小车照片说明已有接口条件，不用来暗示实车有效性；论文也不宣称优于未完整复现的所有主流方法。')

    s = new('结论：语义的价值在于有条件地改变观察决策', '18  总结与答辩',
        '总结三个要点：第一，四接口把类别信念、几何反馈、主动规划和返航执行连起来，但表面始终来自实际观测。第二，原确认和最后的新布局补证都给出条件性类别收益，错误先验完整策略也能改善结果。第三，所有持平、失败和负收益被保留，跨设施共享并未被证明。最终形成可复现代码、主文、补充证据和答辩材料，后续以投稿和学校评阅要求调整文本，不无限扩展实验。谢谢老师。')
    bullets(s,['系统：四接口 + 规划 / 执行层次 + 独立实测重建。',
               f"证据：原确认 +6.16%；新布局 +{d['new_relative_percent']:.2f}%；完整纠错 +2.62%。",
               '结论：类别收益有条件；额外跨设施共享仍未获得支持。'],y=1.91,size=24,gap=.43)
    add_text(s,'感谢各位老师，敬请批评指正。',.70,5.42,11.97,29,TEAL,True,align='center')
    note(s,'三个百分比来自不同预先说明的比较，不能相加；输出包含可编辑 PPT、同版式 PDF、讲稿与问答。')
    if len(slides) != 19:
        raise AssertionError(len(slides))
    for i,s in enumerate(slides,1):
        # Small furniture is shared, not subject to content-area bounds.
        s['elements'].append(dict(kind='text',text=f'语义—几何主动三维建图    /    {i:02d} / 19',x=.62,y=7.16,w=12.07,h=.20,size=9.5,color=GRAY,bold=False,align='right',line=1.1))
    return slides


def color(s):
    return RGBColor.from_string(s.lstrip('#'))


def pptx_build(slides, output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = '面向工业设施的语义—几何主动三维建图'
    prs.core_properties.subject = '中文毕业论文答辩；原始测量与条件性语义收益'
    prs.core_properties.author = ''
    prs.core_properties.last_modified_by = ''
    media_i = 0
    for i,s in enumerate(slides,1):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        for e in s['elements']:
            x,y,w,h = (Inches(e[k]) for k in ('x','y','w','h'))
            if e['kind'] == 'rect':
                sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h)
                sh.fill.solid(); sh.fill.fore_color.rgb=color(e['color'])
                sh.line.fill.background()
            elif e['kind'] == 'text':
                box=sl.shapes.add_textbox(x,y,w,h)
                tf=box.text_frame
                tf.clear(); tf.word_wrap=False
                tf.margin_top=tf.margin_bottom=tf.margin_left=tf.margin_right=0
                for j,tx in enumerate(e['text'].split('\n')):
                    p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
                    from pptx.enum.text import PP_ALIGN
                    p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[e['align']]
                    p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=Pt(e['size']*e['line'])
                    run=p.add_run(); run.text=tx
                    run.font.name=FONTNAME;run.font.size=Pt(e['size']);run.font.bold=e['bold'];run.font.color.rgb=color(e['color'])
                    rpr=run._r.get_or_add_rPr()
                    for tag in ['a:ea','a:cs']:
                        child=OxmlElement(tag);child.set('typeface',FONTNAME);rpr.append(child)
            else:
                pic=sl.shapes.add_picture(str(ROOT/e['paths']['png']),x,y,w,h)
                if 'svg' in e['paths']:
                    media_i+=1
                    part=Part(PackURI(f'/ppt/media/vector_{media_i:03d}.svg'),'image/svg+xml',prs.part.package,(ROOT/e['paths']['svg']).read_bytes())
                    rid=sl.part.relate_to(part,RT.IMAGE)
                    blip=pic._pic.blipFill.blip
                    exts=OxmlElement('a:extLst');ex=OxmlElement('a:ext')
                    ex.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
                    sv=etree.Element('{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip',nsmap={'asvg':'http://schemas.microsoft.com/office/drawing/2016/SVG/main'})
                    sv.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed',rid)
                    ex.append(sv);exts.append(ex);blip.append(exts)
        sl.notes_slide.notes_text_frame.text=f'{i:02d}｜{s["title"]}\n{s["notes"]}'
    prs.save(output)


def pdf_build(slides, output, temp):
    plt.rcParams.update({'pdf.fonttype':42,'svg.fonttype':'none'})
    base=temp/'layout_base.pdf'
    with PdfPages(base,metadata={'Title':'语义—几何主动三维建图答辩','Author':'','Creator':Path(__file__).name}) as pdf:
        for s in slides:
            fig=plt.figure(figsize=(W,H),facecolor='white')
            ax=fig.add_axes([0,0,1,1]);ax.set_xlim(0,W);ax.set_ylim(H,0);ax.axis('off')
            for e in s['elements']:
                if e['kind']=='rect':
                    ax.add_patch(Rectangle((e['x'],e['y']),e['w'],e['h'],facecolor=e['color'],edgecolor=e['edge'] or 'none',linewidth=.5))
                elif e['kind']=='text':
                    xx=e['x']+({'left':0,'center':e['w']/2,'right':e['w']}[e['align']])
                    ax.text(xx,e['y'],e['text'],ha=e['align'],va='top',fontproperties=FontProperties(fname=str(BOLD if e['bold'] else FONT)),fontsize=e['size'],color=e['color'],linespacing=e['line'])
                elif 'pdf' not in e['paths']:
                    ax.imshow(Image.open(ROOT/e['paths']['png']),extent=(e['x'],e['x']+e['w'],e['y']+e['h'],e['y']),aspect='auto')
            pdf.savefig(fig);plt.close(fig)
    # Matplotlib's multipage PDFs share resource dictionaries. Regenerating
    # their original page content after adding forms can alias those names.
    # Import each complete base page as a Form into a new independent page;
    # scientific figure forms then have isolated source resource dictionaries.
    base_doc=pdfium.PdfDocument(base)
    dest=pdfium.PdfDocument.new()
    keep=[base_doc]
    for i,s in enumerate(slides):
        page=dest.new_page(W*72,H*72)
        base_xo=base_doc.page_as_xobject(i,dest);keep.append(base_xo)
        page.insert_obj(base_xo.as_pageobject())
        for e in s['elements']:
            if e['kind']=='picture' and 'pdf' in e['paths']:
                src=pdfium.PdfDocument(ROOT/e['paths']['pdf']);keep.append(src)
                sp=src[0];sw,sh=sp.get_size();sp.close()
                xo=src.page_as_xobject(0,dest);keep.append(xo)
                obj=xo.as_pageobject()
                scale=min(e['w']*72/sw,e['h']*72/sh)
                tx=e['x']*72+(e['w']*72-sw*scale)/2
                ty=(H-e['y']-e['h'])*72+(e['h']*72-sh*scale)/2
                pdfium.raw.FPDFPageObj_Transform(obj,scale,0,0,scale,tx,ty)
                page.insert_obj(obj)
        page.gen_content();page.close()
    dest.save(output)
    for k in reversed(keep):k.close()
    dest.close()


def verify_ppt(path, slides):
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
    with zipfile.ZipFile(path) as z:
        names=set(z.namelist())
        slidefiles=sorted(n for n in names if n.startswith('ppt/slides/slide') and n.endswith('.xml'))
        notefiles=[n for n in names if n.startswith('ppt/notesSlides/notesSlide') and n.endswith('.xml')]
        if len(slidefiles)!=19 or len(notefiles)!=19:
            raise AssertionError('Slide/notes count mismatch')
        for i,s in enumerate(slides,1):
            root=etree.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
            text='\n'.join(root.xpath('//a:t/text()',namespaces=ns))
            if s['title'].replace('₀','0').replace('₅','5') not in text:
                raise AssertionError(f'Missing native title on slide{i}')
        # Resolve package relationships rather than merely checking image count.
        import posixpath
        checked=0
        for name in names:
            if not name.endswith('.rels'):continue
            base=posixpath.dirname(name).replace('/_rels','')
            if name=='_rels/.rels':base=''
            for r in etree.fromstring(z.read(name)):
                if r.get('TargetMode')=='External':
                    raise AssertionError('Unexpected external relationship')
                t=r.get('Target')
                target=t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.join(base,t))
                if target not in names:raise AssertionError((name,target))
                checked+=1
        return dict(slides=len(slidefiles),notes=len(notefiles),relationships_checked=checked,
                    svg_assets=len([n for n in names if n.endswith('.svg')]),native_text=True,
                    office_rendered=False,inspection='OOXML structure and relationship validation only')


def render(pdf, out):
    previews=out/'preview';previews.mkdir()
    doc=pdfium.PdfDocument(pdf)
    thumbs=[]
    for i in range(len(doc)):
        p=doc[i];im=p.render(scale=1.60).to_pil().convert('RGB')
        im.save(previews/f'slide_{i+1:02d}.png',optimize=True)
        thumb=im.copy();thumb.thumbnail((480,270))
        card=Image.new('RGB',(500,304),'#e3eaf0');card.paste(thumb,((500-thumb.width)//2,4))
        dr=ImageDraw.Draw(card);dr.text((12,278),f'{i+1:02d}',font=ImageFont.truetype(str(FONT),17),fill=NAVY)
        thumbs.append(card);p.close()
    doc.close()
    sheet=Image.new('RGB',(1500,304*math.ceil(len(thumbs)/3)),'#cbd5df')
    for i,im in enumerate(thumbs):sheet.paste(im,((i%3)*500,(i//3)*304))
    sheet.save(out/'contact_sheet.png',optimize=True)
    return dict(pages=len(thumbs),preview_dpi=115.2,preview_source='same-layout PDF rendered by PDFium; not PowerPoint/Office')


def qa_text(d):
    return f'''# 答辩问答提纲

每个回答先给结论，再依据需要补充证据。这里的数字与最终答辩稿同源；不把相关检查点、噪声重复和构型数当作独立布局数。

1. **研究究竟解决什么工业问题？** 设备或工位调整后需要重新获取环境表示；本研究处理单次静态任务内，有限预算下如何选择有价值的观察方向。未实现动态避障、变化检测或跨时地图更新。
2. **语义具体是什么？** 局部实验中的类别线索用于给两种公开设备构型设置有限先验，影响尚有几何歧义时的方向选择。它不是在重建评价中对某类设备额外加分。合成类别线索不等于自然图像网络已验证。
3. **G 不知道去补看，是否故意弱化对照？** 不是。G具有同样的模板、安全图、主动几何诊断、重访、反馈和TSDF后端，只去掉类别先验；因此能检验额外类别信息。
4. **四模块是不是四个新网络？** 不是。保留 OV-SDF、STGHP、RPN-UQ、IGCR 四责任接口和规划/执行层次；当前CPU算法并不将每个接口实现为学习网络。贡献是任务化方法及对应机制证据，不靠模块名称制造创新。
5. **与SWAP、SPP、VISTA等有何区别？** 这些方法分别强调语义/重建缺陷引导、语义结构模式预测、查询相关性与视角多样性等。本文聚焦有限配置假设下的类别提前选向、可纠正先验与付费诊断/返航闭环。机制适配不等于完整原方法排行榜；具体边界见贡献对照表。
6. **为什么不是“多看就一定更好”？** 方向、遮挡、采样和剩余预算都会影响边际收益。40个检查点存在短时负差；较大预算下S相对G也转为负差。预测曝光只是规划代理，不是实际质量保证。
7. **为什么用覆盖乘以F1？** 乘积要求二者同时较好，避免只看通道或只盯一台设备；同时报告C、P、R和F1使折中可解释。乘积指标本身不作为原创算法。局部ROI和场景全网格指标不能直接混合。
8. **公开模板、导航图或真值会泄漏答案吗？** G和S都可访问相同双模板和安全图，但不知道实际构型；真值用于终点离线评价，没有反馈给规划器。应承认这些公开先验限制了未知环境泛化范围。
9. **原确认 +6.16% 能代表什么？** 两个既有父布局、两个构型的四组配对，2胜2平；该数字是平均J的相对差，不是提升6.16个百分点，也不是大型独立场景显著性结论。
10. **新布局补证结果如何？** 8条全部合格，4组配对1胜3平；平均J从{d['means']['G']:.6f}到{d['means']['S']:.6f}，相对+{d['new_relative_percent']:.4f}%。它新增两种同族背景/导航布局，双模板仍公开，不是新类别泛化；额外8个预留名额没有使用。
11. **+2.62% 是否证明单个纠错公式有效？** 是完整X/Xnf策略差异。Xnf同时关闭实测更新与未来诊断预测，不能隔离单一公式。两次相同保存状态的后验干预支持下一动作变化，未提供反事实终点质量。
12. **32条失败会不会被排除了？** 128次扩展中96合格、32个B30失败；失败分母与记录保留，失败不冒充已完成的零分样本或从清单中消失。失败属于当前共同前缀后规划不可行，不是物理不可能证明。
13. **36次场景实验为什么没有共享收益？** 局部S/G比较类别有无；场景S/B比较已启用类别后的额外共享。三版三布局的9组S/B动作与质量相同，故共享作用尚未得到支持，但这不推翻局部类别信息的作用。
14. **36次中是不是全流程全部成功？** 原版11合格、1条CELL/G运动完成但原表面评价失败；另外两版各12合格。共36运动完成、35原流程合格。对微小退化面做统一派生评价提供质量，但不把原失败改写为成功。
15. **工位G优于NBV是否证明诊断观测的独立收益？** 不能这样归因。两者在付费诊断分支后路线改变，但首次诊断完成没有新的反馈更新；终点差是整条后续策略的差，不能归于单次观测。
16. **是否验证了SLAM定位精度和真实机器人？** 没有。当前使用准确位姿、CPU传感仿真和实际深度TSDF融合，主要验证主动观测逻辑与重建质量。照片展示已有ROS1/ZED/目标导航接口；实车是后续迁移补充。
17. **研究的创新性边界是什么？** 可辩护之处是有限预算设施建图中的可纠正类别条件观测方法、付费诊断与方向价值解释、完整机制/表面/边界证据。标准贝叶斯理论、四模块命名或改变项目名称不构成单独的新颖性证明。
18. **接下来是否还要一直做实验？** 本轮8条独立布局验证已经结束，预留不自动使用。后续按投稿和学校格式完善表达，只有确实阻断复现或核实的技术问题才做有界修复；不以每场必胜、显著性或录用作为技术收尾条件。
'''


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output',type=Path,default=ROOT/'docs/thesis/defense_20260929')
    args=ap.parse_args()
    out=args.output.resolve()
    if out.exists() and any(out.iterdir()):
        raise SystemExit(f'Refusing to overwrite existing defense output: {out}')
    out.mkdir(parents=True,exist_ok=True)
    d=collect_data();slides=make_slides(d)
    source(__file__)
    source('docs/research/FINAL_CONTRIBUTION_EVIDENCE_MATRIX_20260929.md')
    source('docs/research/FINAL_PAPER_THESIS_TARGETS_20260929.md')
    # Pin the actual source behind mechanism/budget/checkpoint prose as well as the figures.
    for p in ['docs/thesis/ARTICLE_SUBMISSION_EN_20260928.md',
              'audit_results/v35_semantic_chain_review_20260918/result.json']:
        source(p)
    (out/'layout.json').write_text(json.dumps({'schema':'final.defense.layout.v1','slide_inches':[W,H],'slides':slides},ensure_ascii=False,indent=2)+'\n')
    (out/'speaker_notes.md').write_text('# 逐页答辩讲稿\n\n建议总时长约15–18分钟；数字均取保存结果。可编辑PPT内同步保存逐页备注。\n\n'+'\n\n'.join(f'## {i:02d}｜{s["title"]}\n\n{s["notes"]}' for i,s in enumerate(slides,1))+'\n')
    (out/'qa_outline.md').write_text(qa_text(d))
    ppt=out/'semantic_geometric_mapping_defense.pptx'
    pdf=out/'semantic_geometric_mapping_defense.pdf'
    pptx_build(slides,ppt)
    with tempfile.TemporaryDirectory(prefix='nso-defense-layout-') as tmp:
        pdf_build(slides,pdf,Path(tmp))
    checks=verify_ppt(ppt,slides);checks.update(render(pdf,out))
    checks.update(school_degree_major='left blank as requested',physical_results_claimed=False,
                  new_online_tasks=0,new_surface_evaluations=0,new_fusions=0)
    (out/'validation.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n')
    (out/'plotted_data.json').write_text(json.dumps({'new_validation':{k:d[k] for k in ['means','pairs','new_relative_percent']},'source_summary':d['final']},ensure_ascii=False,indent=2)+'\n')
    (out/'README.md').write_text('''# 中文毕业论文答辩材料

19页、16:9。PPTX正文可编辑，逐页备注随文件保存；SVG科学图带PNG兼容回退。PDF从同一layout.json版式构建，已有科学图以原PDF矢量对象嵌入，照片使用完整原始PNG。学校、学位、专业及个人姓名未推测填写。

- `semantic_geometric_mapping_defense.pptx`：可编辑答辩稿，建议安装Noto Sans CJK SC。
- `semantic_geometric_mapping_defense.pdf`：与PPT共用版式数据的展示PDF。
- `preview/slide_*.png`、`contact_sheet.png`：PDFium实际渲染预览。
- `speaker_notes.md` / `qa_outline.md`：逐页讲稿 / 18项问答。
- `layout.json`、`plotted_data.json`、`manifest.json`：版式、数字及输入输出绑定。

本环境无Microsoft Office或LibreOffice。已检查PPT包内19张幻灯片、19份备注、原生文字、SVG和全部关系目标；**没有声称PPT经过Office渲染**。PDF与PNG是同源独立排版的实际渲染。正式答辩前可在学校使用的软件中打开PPT检查字体；需要稳定显示时可直接使用PDF。

新布局8条结果全部纳入：1胜3平，约+2.98%；与原确认+6.16%及完整纠错+2.62%分开。36场景级任务保留原流程35合格/1失败身份和9组S/B持平。照片不作为实车性能证据。

重建到一个全新目录（原文件不覆盖）：
```bash
PYTHONPATH=tmp/final-writing-python:tmp/thesis-python_20260928/site-packages python3 scripts/build_final_defense_20260929.py --output /tmp/nso-defense-rebuild
```

所有输出仅复用保存数据与现有图，不运行World、TSDF、评价或策略。
''')
    (out/'source.py').write_bytes(Path(__file__).read_bytes())
    manifest={'schema':'final.defense.delivery.v1','date':'2026-09-29','builder_sha256':sha(__file__),
              'inputs':[{'path':str(p.relative_to(ROOT)) if p.is_relative_to(ROOT) else str(p),'sha256':sha(p),'bytes':p.stat().st_size} for p in sorted(INPUTS)],
              'outputs':[{'path':str(p.relative_to(out)),'sha256':sha(p),'bytes':p.stat().st_size} for p in sorted(out.rglob('*')) if p.is_file()],
              'validation':checks,'photo_edited':False,'office_rendered':False}
    (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'output':str(out),'slides':len(slides),'new_relative_percent':d['new_relative_percent'],'validation':checks},ensure_ascii=False))


if __name__=='__main__':
    main()
